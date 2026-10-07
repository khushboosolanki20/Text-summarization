"""
Shared fixtures. Test documents are generated in code with PyMuPDF and
python-docx so no binary files need to be committed.
"""

from io import BytesIO

import docx
import pymupdf
import pytest

SAMPLE_TEXT = (
    "Automatic text summarization condenses a long document into a shorter version. "
    "Extractive methods select the most important sentences from the original text. "
    "Abstractive methods generate new sentences that paraphrase the source. "
    "Transformer models such as BART have greatly improved abstractive summarization. "
    "However, these models can only read a limited number of tokens at once."
)


def make_pdf(pages: list[str]) -> bytes:
    """Create a PDF with one text page per list entry."""
    doc = pymupdf.open()
    for text in pages:
        page = doc.new_page()
        page.insert_textbox(pymupdf.Rect(50, 50, 550, 800), text, fontsize=10)
    data = doc.tobytes()
    doc.close()
    return data


def _png_bytes() -> bytes:
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 60, 60), False)
    pix.clear_with(180)
    return pix.tobytes("png")


def make_scanned_pdf(n_pages: int = 2, text_pages: list[str] | None = None) -> bytes:
    """Image-only pages (simulating a scan), optionally followed by text pages."""
    doc = pymupdf.open()
    png = _png_bytes()
    for _ in range(n_pages):
        doc.new_page().insert_image(pymupdf.Rect(50, 50, 500, 700), stream=png)
    for text in text_pages or []:
        doc.new_page().insert_textbox(pymupdf.Rect(50, 50, 550, 800), text, fontsize=10)
    data = doc.tobytes()
    doc.close()
    return data


def make_docx(paragraphs: list[str], with_table: bool = False) -> bytes:
    document = docx.Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if with_table:
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Quarter"
        table.cell(0, 1).text = "Revenue"
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def sample_text() -> str:
    return SAMPLE_TEXT
