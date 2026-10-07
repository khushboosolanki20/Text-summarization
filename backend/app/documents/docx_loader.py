"""
DOCX loader built on python-docx.

Body paragraphs are extracted in order. Tables are not included in the text
because their cells are fragments ("Q3", "12.4%") rather than sentences and
would confuse sentence-based summarization; the user is warned instead.
"""

from io import BytesIO

import docx

from app.documents.base import ExtractedDocument
from app.errors import CorruptedFileError, EmptyDocumentError


def load_docx(data: bytes, filename: str = "document.docx") -> ExtractedDocument:
    try:
        document = docx.Document(BytesIO(data))
    except Exception as exc:  # BadZipFile, PackageNotFoundError, KeyError, XML errors ...
        raise CorruptedFileError("The DOCX file is corrupted or invalid and could not be opened.") from exc

    paragraphs = []
    for paragraph in document.paragraphs:
        # Soft line breaks (Shift+Enter) arrive as "\n" inside a paragraph.
        text = paragraph.text.replace("\n", " ").strip()
        if text:
            paragraphs.append(text)

    if not paragraphs:
        raise EmptyDocumentError("No text paragraphs were found in the DOCX file.")

    warnings = []
    if document.tables:
        warnings.append(f"{len(document.tables)} table(s) were not included in the summary input.")

    props = document.core_properties
    metadata = {k: v for k, v in {"title": props.title, "author": props.author}.items() if v}
    text = "\n\n".join(paragraphs)

    return ExtractedDocument(
        text=text,
        file_type="docx",
        filename=filename,
        pages=[text],
        warnings=warnings,
        metadata={**metadata, "paragraph_count": len(paragraphs)},
    )
